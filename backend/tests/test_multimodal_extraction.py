import os
import pytest
from app.services.pdf_extractor import extract_document_text, SUPPORTED_EXTENSIONS

def test_supported_extensions_registered():
 assert ".pdf" in SUPPORTED_EXTENSIONS
 assert ".png" in SUPPORTED_EXTENSIONS
 assert ".jpg" in SUPPORTED_EXTENSIONS
 assert ".jpeg" in SUPPORTED_EXTENSIONS
 assert ".docx" in SUPPORTED_EXTENSIONS
 assert ".txt" in SUPPORTED_EXTENSIONS

def test_extract_plain_text_file(tmp_path):
 txt_file = tmp_path / "sample_contract.txt"
 txt_file.write_text("THIS IS A FREELANCE SERVICES AGREEMENT BETWEEN PARTY A AND PARTY B.")
 
 result = extract_document_text(str(txt_file))
 assert result["text"].strip() == "THIS IS A FREELANCE SERVICES AGREEMENT BETWEEN PARTY A AND PARTY B."
 assert result["strategy"] == "plain_text"
 assert result["is_scanned"] is False

def test_extract_docx_file(tmp_path):
    import docx
    docx_file = tmp_path / "sample_contract.docx"
    doc = docx.Document()
    doc.add_heading("NON-DISCLOSURE AGREEMENT", level=1)
    doc.add_paragraph("This NDA is entered into on August 13, 2026.")
    doc.save(str(docx_file))

    result = extract_document_text(str(docx_file))
    assert "NON-DISCLOSURE AGREEMENT" in result["text"]
    assert "August 13, 2026" in result["text"]
    assert result["strategy"] == "python_docx"


def test_preprocess_image_downscaling():
    from PIL import Image
    from app.services.pdf_extractor import _preprocess_image_for_ocr

    # Create a 3000x4000 image
    large_img = Image.new("RGB", (3000, 4000), color="white")
    processed = _preprocess_image_for_ocr(large_img, max_dim=2048)
    
    assert max(processed.size) <= 2048
    assert processed.size == (1536, 2048)
    assert processed.mode == "RGB"


def test_preprocess_image_rgba_to_rgb():
    from PIL import Image
    from app.services.pdf_extractor import _preprocess_image_for_ocr

    rgba_img = Image.new("RGBA", (500, 500), color=(255, 0, 0, 128))
    processed = _preprocess_image_for_ocr(rgba_img)
    assert processed.mode == "RGB"


def test_ocr_image_with_gemini_fallback(monkeypatch):
    from unittest.mock import MagicMock
    from PIL import Image
    import google.generativeai as genai
    from app.services.pdf_extractor import ocr_image_with_gemini

    calls = []

    def mock_generative_model(model_name):
        calls.append(model_name)
        mock_model = MagicMock()
        if len(calls) == 1:
            # First model simulates 429 quota exhaustion
            mock_model.generate_content.side_effect = Exception("429 ResourceExhausted quota exceeded")
        else:
            # Fallback model succeeds
            mock_resp = MagicMock()
            mock_resp.text = "AGREEMENT BETWEEN CLIENT AND VENDOR"
            mock_model.generate_content.return_value = mock_resp
        return mock_model

    monkeypatch.setattr(genai, "GenerativeModel", mock_generative_model)
    monkeypatch.setattr("app.services.pdf_extractor._ensure_genai_configured", lambda: None)

    img = Image.new("RGB", (200, 200), color="white")
    text = ocr_image_with_gemini(img)

    assert text == "AGREEMENT BETWEEN CLIENT AND VENDOR"
    assert len(calls) >= 2  # Proves failover cascade executed


def test_extract_image_text_missing_file():
    from fastapi import HTTPException
    from app.services.pdf_extractor import extract_image_text

    with pytest.raises(HTTPException) as exc_info:
        extract_image_text("/nonexistent/contract_scan.jpg")
    assert exc_info.value.status_code == 404


def test_extract_image_text_success(tmp_path, monkeypatch):
    from PIL import Image
    from app.services.pdf_extractor import extract_image_text

    img_path = tmp_path / "test_scan.png"
    Image.new("RGB", (100, 100), color="white").save(str(img_path))

    monkeypatch.setattr(
        "app.services.pdf_extractor.ocr_image_with_gemini",
        lambda img: "EMPLOYMENT BOND AGREEMENT 2026"
    )

    res = extract_image_text(str(img_path))
    assert res["text"] == "EMPLOYMENT BOND AGREEMENT 2026"
    assert res["page_count"] == 1
    assert res["is_scanned"] is True
    assert res["strategy"] == "gemini_vision_ocr"

